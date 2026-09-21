from datetime import date
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

out_dir = r"D:\mujahidCXN\brotherLPG\backend\docs"
os.makedirs(out_dir, exist_ok=True)
out_path = os.path.join(out_dir, "Brother_LPG_Client_User_Guide.docx")

doc = Document()

for section in doc.sections:
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.9)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)


def set_run_font(run, size=11, bold=False, color=None, name="Calibri"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_heading_styled(text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(
            run,
            size=18 if level == 1 else 14,
            bold=True,
            color=(0x1F, 0x4E, 0x79),
        )
    return p


def add_para(text, size=11, bold=False, space_after=8):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    return p


def add_hyperlink(paragraph, text, url, size=10):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")

    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rPr.append(color)

    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)

    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(size * 2))
    rPr.append(sz)

    szCs = OxmlElement("w:szCs")
    szCs.set(qn("w:val"), str(size * 2))
    rPr.append(szCs)

    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), "Calibri")
    rFonts.set(qn("w:hAnsi"), "Calibri")
    rPr.append(rFonts)

    new_run.append(rPr)
    text_elem = OxmlElement("w:t")
    text_elem.text = text
    new_run.append(text_elem)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_link_para(label, url):
    p = doc.add_paragraph()
    label_run = p.add_run(f"{label}: ")
    set_run_font(label_run, size=11, bold=True, color=(0x2E, 0x7D, 0x32))
    add_hyperlink(p, url, url, size=10)
    p.paragraph_format.space_after = Pt(10)
    return p


def add_hr():
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1F4E79")
    pBdr.append(bottom)
    pPr.append(pBdr)


title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("Brother LPG")
set_run_font(r, size=28, bold=True, color=(0x1F, 0x4E, 0x79))
title.paragraph_format.space_after = Pt(4)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = subtitle.add_run("Client User Guide & Training Videos")
set_run_font(r, size=16, bold=True, color=(0x2E, 0x75, 0xB6))
subtitle.paragraph_format.space_after = Pt(6)

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = meta.add_run(f"Version 1.0  |  {date.today().strftime('%d %B %Y')}")
set_run_font(r, size=11, color=(0x66, 0x66, 0x66))
meta.paragraph_format.space_after = Pt(18)

add_hr()

add_heading_styled("1. Introduction", 1)
add_para(
    "This document is a simple client guide for the Brother LPG management system. "
    "It explains each main module in short, everyday language and provides a training video link "
    "so your team can learn by watching the actual screens."
)
add_para(
    "How to use this guide: read the short description for the module you need, then open the video link "
    "and follow the same steps in your system."
)

add_heading_styled("2. Training Video Index", 1)
add_para(
    "All training videos are listed below by module. Click any link to open the video in Google Drive."
)

modules = [
    {
        "no": "1.1",
        "title": "Employee & User",
        "url": "https://drive.google.com/file/d/16ldaJNz84ApzXBEQclkbATusF2kxP2fw/view?usp=drive_link",
        "desc": (
            "This module covers creating and managing employees, and linking them with system login users. "
            "You can add employee details (department, status, contact info), create user accounts for staff, "
            "and keep employee and user records organized for daily operations."
        ),
        "points": [
            "Add / edit employee profiles",
            "Create login users for staff",
            "Set employment status (Active / Inactive / Terminated)",
            "Connect employee records with system access",
        ],
    },
    {
        "no": "1.2",
        "title": "Users, Roles & Permissions",
        "url": "https://drive.google.com/file/d/1KyUjt_5vZ9UafsQojwDqKgf9Y26LeHG6/view?usp=drive_link",
        "desc": (
            "This module controls who can access what inside the system. "
            "You can manage roles (for example Admin, Sales, Accounts), assign permissions to each role, "
            "and make sure every user only sees the screens and actions allowed for their job."
        ),
        "points": [
            "Create and update user accounts",
            "Define roles and permission sets",
            "Assign roles to users",
            "Restrict modules by access rights for better security",
        ],
    },
    {
        "no": "1.3",
        "title": "Supplier, Storage Tank & LPG Receipt",
        "url": "https://drive.google.com/file/d/1JkqbDrkfyvsNVROzweCxwKVRDVZxDYV9/view?usp=drive_link",
        "desc": (
            "This module handles the purchase side of LPG operations. "
            "You can register suppliers, manage storage tank details and status, "
            "and record LPG receipts when gas is received into your tanks from suppliers."
        ),
        "points": [
            "Add and manage supplier records",
            "Maintain storage tank capacity and status",
            "Record LPG receipt / inward entries",
            "Track supplier-linked stock intake",
        ],
    },
    {
        "no": "1.4",
        "title": "Supplier Payment & Accounts",
        "url": "https://drive.google.com/file/d/1Kw8jVAa6VWgxgOt_BRlgmImkeOs-chWV/view?usp=drive_link",
        "desc": (
            "This module is for paying suppliers and managing finance accounts. "
            "You can record supplier payments (cash, bank, cheque, online), update payable balances, "
            "and maintain account records used across business transactions."
        ),
        "points": [
            "Record supplier payments",
            "Choose payment method (Cash / Bank / Cheque / Online)",
            "Manage accounts (Cash, Bank, Payable, Receivable, etc.)",
            "Keep supplier dues and payment history clear",
        ],
    },
    {
        "no": "1.5",
        "title": "Customer & Sale",
        "url": "https://drive.google.com/file/d/16DXZhsAeOgPRCEoKb2ryTfHGQ_9QAykF/view?usp=drive_link",
        "desc": (
            "This module covers customer master data and creating sales. "
            "You can add customers, create cash or credit sales, select items (cylinders / LPG), "
            "and confirm invoices for dispatch and billing."
        ),
        "points": [
            "Create and manage customer profiles",
            "Create sales invoices (Cash / Credit)",
            "Add sale items and quantities",
            "Confirm and track sale status",
        ],
    },
    {
        "no": "1.6",
        "title": "Customer Sale Payment",
        "url": "https://drive.google.com/file/d/1Fz31KL8EXNRSRdN019JoBIZ3ouF1VTtf/view?usp=drive_link",
        "desc": (
            "This module is for collecting money from customers against sales. "
            "You can receive full or partial payments, update payment status "
            "(Unpaid / Partial / Paid), and keep customer receivable balances up to date."
        ),
        "points": [
            "Receive customer payments against invoices",
            "Support full and partial collection",
            "Track payment status of each sale",
            "Update customer receivable balance",
        ],
    },
    {
        "no": "1.7",
        "title": "Sale Return",
        "url": "https://drive.google.com/file/d/11ItW2o20JEcVE33iDYPnLG4vlcFjRSXt/view?usp=drive_link",
        "desc": (
            "This module handles returns after a sale is completed. "
            "You can return defective or wrong items, select a return reason "
            "(for example defective valve, damaged in transit, customer request), "
            "and adjust the related sale and customer ledger."
        ),
        "points": [
            "Create sale return entries",
            "Select return reason and quantity",
            "Update sale status (Partially Returned / Returned)",
            "Credit customer ledger where applicable",
        ],
    },
    {
        "no": "1.8",
        "title": "Customer Refund Payment",
        "url": "https://drive.google.com/file/d/182pVc0mj7lo6UyuiJ0GGIC6c-hXNodft/view?usp=drive_link",
        "desc": (
            "This module is used when money needs to be returned to the customer "
            "(for example after a return or overpayment). "
            "You can process refund payments and clear refund-due amounts against the customer account."
        ),
        "points": [
            "Process customer refund payments",
            "Link refund with return / due amount",
            "Update payment and ledger status",
            "Maintain clear customer refund history",
        ],
    },
    {
        "no": "1.9",
        "title": "Dashboard & Refilling Batch",
        "url": "https://drive.google.com/file/d/1s5xf7-fWxDCTVsqCJYc3pAFc9fiMuhBF/view?usp=drive_link",
        "desc": (
            "This module shows business overview on the dashboard and manages cylinder refilling batches. "
            "The dashboard helps you quickly see key numbers, while filling batches record production / "
            "refilling work against available LPG stock and cylinder inventory."
        ),
        "points": [
            "View dashboard summary for quick monitoring",
            "Create and manage refilling / filling batches",
            "Track filled vs empty cylinder movement",
            "Connect plant operations with stock status",
        ],
    },
]

for m in modules:
    p = doc.add_paragraph()
    r = p.add_run(f"{m['no']}  {m['title']}")
    set_run_font(r, size=11, bold=True)
    p.paragraph_format.space_after = Pt(2)
    add_link_para("Watch video", m["url"])

add_hr()
add_heading_styled("3. Module Details", 1)
add_para(
    "Below is a short description of each module. Use these notes with the matching training video."
)

for m in modules:
    add_heading_styled(f"{m['no']}  {m['title']}", 2)
    add_para(m["desc"])
    add_para("What you can do:", bold=True, space_after=4)
    for point in m["points"]:
        bp = doc.add_paragraph(style="List Bullet")
        run = bp.add_run(point)
        set_run_font(run, size=11)
        bp.paragraph_format.space_after = Pt(2)
    add_link_para("Training video", m["url"])
    add_hr()

add_heading_styled("4. Quick Tips for Client Team", 1)
tips = [
    "Always create master data first (Employees, Users, Suppliers, Customers, Accounts, Tanks) before daily transactions.",
    "Use Roles & Permissions carefully so staff only access the modules needed for their work.",
    "Record LPG receipts before creating filling batches, so stock remains accurate.",
    "For credit sales, regularly collect customer payments to keep receivables under control.",
    "Use Sale Return and Refund modules only when needed, and select the correct reason for clear records.",
    "If a video link asks for Google sign-in, open it with the Google account that has access to the shared Drive folder.",
]
for tip in tips:
    bp = doc.add_paragraph(style="List Bullet")
    run = bp.add_run(tip)
    set_run_font(run, size=11)

add_heading_styled("5. Support", 1)
add_para(
    "If you face any issue while using a module, note the screen name and what you were trying to do, "
    "then contact your Brother LPG support team with that detail. You can also re-watch the related video from Section 2."
)

footer = doc.add_paragraph()
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.paragraph_format.space_before = Pt(24)
r = footer.add_run("— End of Client User Guide —")
set_run_font(r, size=10, color=(0x88, 0x88, 0x88))

doc.save(out_path)
print(out_path)
print("OK")
